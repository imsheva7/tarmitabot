"""Inspect generated XLSX values, types and headers, without starting a bot."""
from io import BytesIO
from openpyxl import load_workbook
import pytest
from erpbot.reporting.models import Sheet,WorkbookReport
from erpbot.reporting.xlsx import build_workbook
from erpbot.services import reports
from erpbot.reporting.catalogue import ADMIN_EXPORTS


@pytest.mark.parametrize('value',['=1+1','+SUM(1,2)','@SUM(1,2)','-1+1','=HYPERLINK("https://example.invalid","click")'])
def test_untrusted_text_remains_literal_and_never_becomes_formula(value):
    report=WorkbookReport('test.xlsx',(Sheet('گزارش',('عنوان',),((value,),)),))
    workbook=load_workbook(BytesIO(build_workbook(report)),data_only=False)
    try:
        cell=workbook.active['A2']
        assert cell.value==value and cell.data_type=='s'
        assert workbook.active.sheet_view.rightToLeft
        assert workbook.active.freeze_panes=='A2'
    finally: workbook.close()


def test_workbook_preserves_types_unicode_and_leading_zeroes():
    report=WorkbookReport('test.xlsx',(Sheet('داده',('نام','تلفن','مقدار','خالی'),(('کاربر فارسی','09123456789',1234,None),)),))
    workbook=load_workbook(BytesIO(build_workbook(report)))
    try:
        assert tuple(c.value for c in workbook.active[2])==('کاربر فارسی','09123456789',1234,None)
        assert workbook.active['B2'].data_type=='s' and workbook.active['C2'].data_type=='n'
    finally: workbook.close()


@pytest.mark.parametrize('key',list(ADMIN_EXPORTS))
def test_actual_admin_workbooks_have_same_filename_sheet_and_columns(assets_db,key):
    with assets_db.connection() as conn: report=reports.admin_export(conn,101,key)
    workbook=load_workbook(BytesIO(build_workbook(report)))
    try:
        _,filename,title,mapping=ADMIN_EXPORTS[key]
        assert report.filename==filename and workbook.sheetnames==[title]
        assert tuple(c.value for c in workbook.active[1])==tuple(mapping.values())
    finally: workbook.close()


def test_empty_personal_report_has_two_sheets_with_headers(assets_db):
    with assets_db.connection() as conn: report=reports.personal_export(conn,202)
    workbook=load_workbook(BytesIO(build_workbook(report)))
    try:
        assert workbook.sheetnames==['درخواست‌های من','اموال من']
        assert [sheet.max_row for sheet in workbook.worksheets]==[1,1]
    finally: workbook.close()
