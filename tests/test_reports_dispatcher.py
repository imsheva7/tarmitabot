"""Actual report Router, document bytes and existing approval/cancel buttons, offline."""
import os
from pathlib import Path
import subprocess
import sys


def test_reports_with_real_dispatcher_and_generated_workbooks(tmp_path):
    script=r'''
import asyncio
from io import BytesIO
from pathlib import Path
import jdatetime
from openpyxl import load_workbook
from aiogram.client.session.base import BaseSession
from aiogram.methods import AnswerCallbackQuery,EditMessageText,SendMessage,SendDocument
from aiogram.types import Message,Update,User
from cryptography.fernet import Fernet
from erpbot.bootstrap import initialize_runtime
from erpbot.core import runtime
from erpbot.core.settings import Settings
from erpbot.loader import register_handlers

class OfflineSession(BaseSession):
    def __init__(self):
        super().__init__()
        self.calls=[]
    async def close(self): pass
    async def make_request(self,bot,method,timeout=None):
        self.calls.append(method)
        if isinstance(method,AnswerCallbackQuery): return True
        assert isinstance(method,(SendMessage,EditMessageText,SendDocument)),type(method)
        return Message(message_id=100+len(self.calls),date=0,chat={'id':method.chat_id,'type':'private'},text=getattr(method,'text','Document'))
    async def stream_content(self,*args,**kwargs):
        raise AssertionError('No network')
        yield b''

async def main():
    app=initialize_runtime(Settings.from_env({
        'BOT_TOKEN':'123456789:ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijk',
        'FERNET_KEY':Fernet.generate_key().decode(),'GROUP_ID':'-100999',
    },data_directory=Path.cwd()),configure_logging=False)
    runtime.init_db()
    runtime.get_current_tehran_time=lambda:jdatetime.datetime(1405,7,24,12,30)
    with runtime.get_connection() as conn:
        conn.executemany('INSERT INTO staff_mapping(staff_name,telegram_user_id,role,leave_balance) VALUES(?,?,?,?)',[
            ('Admin',101,'main_admin',5),('Employee <test>',202,'staff',5),('Other',303,'staff',5),
            ('Finance',404,'finance',5),('Purchase',505,'purchase',5)])
        for uid,title in ((202,'=1+1'),(303,'Other confidential request')):
            rid=conn.execute("INSERT INTO requests(telegram_user_id,user_name,request_type,col1,col2,col3,status,created_at,month_tag,sensitive_data) VALUES(?,?,'ماموریت',?,'شرح','جزئیات','⏳ در انتظار تأیید نهایی مدیریت','1405/07/24 12:00','1405/07',?)",(uid,'Employee' if uid==202 else 'Other',title,runtime.encrypt_data('hidden-request-secret'))).lastrowid
            conn.execute('UPDATE requests SET req_code=? WHERE id=?',(f'REQ-1405-{rid:05d}',rid))
            sid=conn.execute('SELECT staff_id FROM staff_mapping WHERE telegram_user_id=?',(uid,)).fetchone()[0]
            conn.execute("INSERT INTO user_assets(staff_id,staff_name,asset_type,details,encrypted_password) VALUES(?,?,'کارت سوخت',?,?)",(sid,'Owner','Own asset' if uid==202 else 'Other confidential asset',runtime.encrypt_data('hidden-asset-secret')))
        conn.execute("INSERT INTO requests(telegram_user_id,user_name,request_type,col1,col2,location,status,month_tag,start_date,end_date) VALUES(202,'Employee <test>','مرخصی','نوع: استحقاقی','شرح',2,'✅ تأیید شده قطعی','1405/07','1405/07/23','1405/07/24')")
    register_handlers()
    session=OfflineSession()
    app.bot.session=session
    app.bot._me=User(id=123456789,is_bot=True,first_name='Bot',username='test_bot')
    sequence=0
    async def click(data,user=202):
        nonlocal sequence
        sequence+=1
        session.calls.clear()
        query={'id':str(sequence),'from':{'id':user,'is_bot':False,'first_name':'Test'},
            'chat_instance':'test','data':data,'message':{'message_id':10,'date':0,'chat':{'id':user,'type':'private'},
            'from':{'id':123456789,'is_bot':True,'first_name':'Bot'},'text':'Menu'}}
        await app.dispatcher.feed_update(app.bot,Update.model_validate({'update_id':sequence,'callback_query':query}))
    async def today(user=101,group=False):
        nonlocal sequence
        sequence+=1
        session.calls.clear()
        message={'message_id':sequence,'date':0,'chat':{'id':-100999 if group else user,'type':'supergroup' if group else 'private'},
            'from':{'id':user,'is_bot':False,'first_name':'Test'},'text':'/today',
            'entities':[{'type':'bot_command','offset':0,'length':6}]}
        await app.dispatcher.feed_update(app.bot,Update.model_validate({'update_id':sequence,'message':message}))
    def document():
        call=next(call for call in session.calls if isinstance(call,SendDocument))
        return load_workbook(BytesIO(call.document.data),data_only=False)
    await click('menu_reports_tracking')
    assert isinstance(session.calls[0],EditMessageText)
    await click('my_stats')
    assert 'Employee &lt;test&gt;' in session.calls[0].text and '2.0 روز' in session.calls[0].text
    await click('my_requests')
    assert all('Other confidential' not in (getattr(call,'text','') or '') for call in session.calls)
    pending=next(call for call in session.calls if isinstance(call,SendMessage) and call.reply_markup)
    assert pending.reply_markup.inline_keyboard[0][0].callback_data=='cancel_REQ-1405-00001'
    await click('my_excel_report')
    workbook=document()
    assert workbook.sheetnames==['درخواست‌های من','اموال من']
    values=[cell.value for sheet in workbook for row in sheet for cell in row]
    assert 'Other confidential request' not in values and 'Other confidential asset' not in values
    assert not any('hidden-' in str(value) for value in values)
    assert workbook.worksheets[0]['C2'].value=='=1+1' and workbook.worksheets[0]['C2'].data_type=='s'
    workbook.close()
    await click('cancel_REQ-1405-00001')
    with runtime.get_connection() as conn:
        assert 'لغو شده' in conn.execute('SELECT status FROM requests WHERE id=1').fetchone()[0]
    for actor in (101,404,505):
        await click('adm_cat_excel',actor)
        assert isinstance(session.calls[0],EditMessageText)
        for key in ('req','staff','assets','ins'):
            await click('xcl_'+key,actor)
            workbook=document()
            assert len(workbook.sheetnames)==1 and workbook.active.max_row>=1
            workbook.close()
    await click('xcl_staff',202)
    assert not any(isinstance(call,SendDocument) for call in session.calls) and session.calls[-1].show_alert
    await click('xcl_invalid',101)
    assert session.calls[-1].show_alert
    await click('admin_pending_reqs',101)
    pending=next(call for call in session.calls if isinstance(call,SendMessage) and call.reply_markup)
    assert pending.reply_markup.inline_keyboard[0][0].callback_data=='approve_2'
    await click('approve_2',101)
    with runtime.get_connection() as conn:
        assert conn.execute('SELECT status FROM requests WHERE id=2').fetchone()[0]=='✅ تأیید شده قطعی'
    await click('admin_list_staff',404)
    assert 'Employee &lt;test&gt;' in session.calls[0].text
    await click('admin_today',101)
    assert 'Employee &lt;test&gt;' in session.calls[0].text
    await today(404)
    assert 'Employee &lt;test&gt;' in session.calls[0].text
    await today(202)
    assert 'غیرمجاز' in session.calls[0].text
    await today(101,group=True)
    assert session.calls==[]
    with runtime.get_connection() as conn:
        conn.execute('UPDATE staff_mapping SET is_active=0 WHERE telegram_user_id=202')
        conn.execute("UPDATE staff_mapping SET role='staff' WHERE telegram_user_id=404")
    await click('my_excel_report')
    assert not any(isinstance(call,SendDocument) for call in session.calls)
    await click('xcl_staff',404)
    assert not any(isinstance(call,SendDocument) for call in session.calls)
    await app.close()

asyncio.run(main())
'''
    root=Path(__file__).resolve().parents[1]
    env=dict(os.environ,PYTHONPATH=str(root),PYTHONIOENCODING='utf-8')
    result=subprocess.run([sys.executable,'-c',script],cwd=tmp_path,env=env,capture_output=True,text=True,encoding='utf-8',timeout=60)
    assert result.returncode==0,result.stdout+result.stderr
